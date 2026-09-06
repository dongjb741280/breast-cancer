"""回填评测结果到 xlsx「真实-实际评测」sheet。

重跑 agent + judge（默认 qwen3.8-max，claude 太慢），把 7 例六项结果写回 Excel。
用法：HH_API_KEY=... HH_MODEL=qwen3.8-max python knowledge_base/fill_xlsx.py
"""
import os
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from agent import Agent
from eval_agent import judge
from extract_case import CASE_JSON, extract_case_text
from gold_standard import SIX, load_real_gold

XLSX = Path("Data_Cleaning/doc/HER2决策系统-病例评测表-V1-虚构与真实分开.xlsx")


def agent_to_six(out):
    """agent 输出 → 六项「系统输出」文本（顺序对应 SIX）。"""
    pop = out.get("人群", "未知")
    stage = out.get("治疗阶段", "未知")
    strat = out.get("分层条件") or ""
    plans = [r["方案"] for r in out.get("方案", [])]
    ev = [f"{r['方案']}[{r['证据类别']} p{r['来源页码']}]" for r in out.get("方案", [])]
    risks = "；".join(out.get("风险提示", [])) or "无"
    need = "；".join(out.get("需补充", [])) or "无"
    advice = out.get("治疗建议") or ("、".join(plans) if plans else "不推荐（资料不足/待核验）")
    return [
        out.get("人群判断") or f"人群={pop}" + ("" if out.get("资料充足") else "（资料不足）"),
        f"阶段={stage}；{strat}".strip("；"),
        advice,
        "；".join(ev) if ev else "无",
        risks,
        need,
    ]


def main():
    if not os.environ.get("HH_API_KEY"):
        sys.exit("缺 HH_API_KEY")
    golds = load_real_gold()
    agent = Agent()
    wb = load_workbook(XLSX)
    ws = wb["真实-实际评测"]

    for ci, (cid, gold) in enumerate(golds.items()):
        out = agent.run(extract_case_text(CASE_JSON[cid]))
        v = judge(cid, gold, out)
        hi = v.get("高风险错误") or []
        lines = agent_to_six(out)
        print(f"{cid}: " + " ".join(f"{k}={v.get(k,'?')}" for k in SIX))
        for ri, item in enumerate(SIX):
            row = 2 + ci * 6 + ri
            score = int(v.get(item, 0))
            ws.cell(row=row, column=6, value=lines[ri])          # 系统输出
            ws.cell(row=row, column=7, value="是" if score == 2 else ("部分" if score == 1 else "否"))  # 是否正确
            ws.cell(row=row, column=11, value="是" if hi else "否")  # 是否触发人工复核
            ws.cell(row=row, column=12, value=(v.get("理由") or "") + (("；高风险：" + "；".join(hi)) if hi else ""))  # 备注
            ws.cell(row=row, column=13, value=score)             # 单项得分
        wb.save(XLSX)  # 增量保存，避免中途超时丢进度

    wb.save(XLSX)
    print(f"\n已回填 {len(golds)} 例 -> {XLSX}")


if __name__ == "__main__":
    main()
